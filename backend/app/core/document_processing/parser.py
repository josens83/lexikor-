"""
Document Parser for various file formats
"""

import os
from typing import Dict, Optional
from pathlib import Path
import pypdf
from docx import Document as DocxDocument


class DocumentParser:
    """Parse documents and extract text"""

    @staticmethod
    def parse(file_path: str) -> Dict[str, any]:
        """
        Parse document and extract content

        Args:
            file_path: Path to document

        Returns:
            Dict containing extracted text and metadata
        """
        extension = Path(file_path).suffix.lower()

        if extension == '.pdf':
            return DocumentParser.parse_pdf(file_path)
        elif extension in ['.docx', '.doc']:
            return DocumentParser.parse_docx(file_path)
        elif extension == '.txt':
            return DocumentParser.parse_txt(file_path)
        elif extension == '.hwp':
            return DocumentParser.parse_hwp(file_path)
        else:
            raise ValueError(f"Unsupported file format: {extension}")

    @staticmethod
    def parse_pdf(file_path: str) -> Dict[str, any]:
        """Parse PDF document"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = pypdf.PdfReader(file)
                num_pages = len(pdf_reader.pages)

                # Extract text from all pages
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"

                metadata = {
                    "num_pages": num_pages,
                    "pdf_info": pdf_reader.metadata if pdf_reader.metadata else {}
                }

                return {
                    "text": text.strip(),
                    "metadata": metadata,
                    "success": True
                }

        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }

    @staticmethod
    def parse_docx(file_path: str) -> Dict[str, any]:
        """Parse DOCX document"""
        try:
            doc = DocxDocument(file_path)

            # Extract text from paragraphs
            text = "\n".join([para.text for para in doc.paragraphs])

            # Extract tables if any
            tables_text = ""
            for table in doc.tables:
                for row in table.rows:
                    tables_text += "\t".join([cell.text for cell in row.cells]) + "\n"

            full_text = text + "\n" + tables_text

            metadata = {
                "num_paragraphs": len(doc.paragraphs),
                "num_tables": len(doc.tables)
            }

            return {
                "text": full_text.strip(),
                "metadata": metadata,
                "success": True
            }

        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }

    @staticmethod
    def parse_txt(file_path: str) -> Dict[str, any]:
        """Parse plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()

            return {
                "text": text.strip(),
                "metadata": {},
                "success": True
            }

        except UnicodeDecodeError:
            # Try different encodings
            for encoding in ['cp949', 'euc-kr', 'latin-1']:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        text = file.read()
                    return {
                        "text": text.strip(),
                        "metadata": {"encoding": encoding},
                        "success": True
                    }
                except:
                    continue

            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": "Could not decode file"
            }

        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }

    @staticmethod
    def parse_hwp(file_path: str) -> Dict[str, any]:
        """Parse HWP (Hangul Word Processor) document"""
        # TODO: Implement HWP parsing
        # Requires olefile or hwp library
        return {
            "text": "[HWP 파일 파싱 미구현 - 추후 구현 예정]",
            "metadata": {},
            "success": False,
            "error": "HWP parsing not implemented yet"
        }


class TextChunker:
    """Chunk text for vector database indexing"""

    @staticmethod
    def chunk_text(
        text: str,
        chunk_size: int = 500,
        chunk_overlap: int = 50
    ) -> list[str]:
        """
        Split text into chunks

        Args:
            text: Text to chunk
            chunk_size: Size of each chunk in characters
            chunk_overlap: Overlap between chunks

        Returns:
            List of text chunks
        """
        chunks = []
        start = 0

        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]
            chunks.append(chunk)
            start = end - chunk_overlap

        return chunks

    @staticmethod
    def semantic_chunk(text: str) -> list[str]:
        """
        Chunk text by semantic boundaries (paragraphs, sections)

        Args:
            text: Text to chunk

        Returns:
            List of semantic chunks
        """
        # Split by paragraphs
        paragraphs = text.split('\n\n')

        # Group small paragraphs
        chunks = []
        current_chunk = ""

        for para in paragraphs:
            if len(current_chunk) + len(para) < 1000:
                current_chunk += para + "\n\n"
            else:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                current_chunk = para + "\n\n"

        if current_chunk:
            chunks.append(current_chunk.strip())

        return chunks
